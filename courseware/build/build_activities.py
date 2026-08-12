#!/usr/bin/env python3
"""Create ten connected Orion activity packs with starter/completed workbooks."""
import math
import os
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import course_data as C
from data_domain1 import DOMAIN1
from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5

ACTIVITIES = DOMAIN1 + DOMAIN2 + DOMAIN3 + DOMAIN4 + DOMAIN5
ROOT = HERE.parents[1]
OUT = ROOT / "activities"
BLUE, TEAL, INK, GREY, LIGHT = "1F6FEB", "10B981", "161B26", "5B6372", "F5F8FC"
THIN = Side(style="thin", color="D9E2F0")


def safe_title(text):
    return "".join(ch for ch in text if ch not in '<>:"/\\|?*').strip()


def setup_sheet(ws, title, subtitle=""):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:H2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color=GREY)
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18


def header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = 30


def finish(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.row > 2:
                cell.font = Font(name="Arial", size=10, bold=cell.font.bold, color=cell.font.color)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = Border(bottom=THIN)
    ws.auto_filter.ref = ws.dimensions if ws.max_row >= 4 else None


def add_sheet(wb, name, title, headers, rows, subtitle=""):
    ws = wb.create_sheet(name[:31])
    setup_sheet(ws, title, subtitle)
    header(ws, 4, headers)
    for r, values in enumerate(rows, 5):
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
    finish(ws)
    return ws


def instructions_sheet(wb, activity):
    ws = wb.active
    ws.title = "Instructions"
    setup_sheet(ws, f"Activity {activity['num']:02d} — {activity['title']}", C.CASE_STUDY["company"])
    entries = [
        ("Mapping", activity["codes"]), ("Objective", activity["objective"]),
        ("Scenario", activity["scenario"]), ("Deliverable", activity["build"]),
        ("Acceptance check", activity["test"]),
        ("Learner note", "Follow the detailed procedure in README.md or the Learner Guide. Blue cells contain examples or formulas; complete the remaining evidence."),
    ]
    for row, (label, value) in enumerate(entries, 4):
        ws.cell(row, 1, label).font = Font(name="Arial", bold=True, color=BLUE)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.cell(row, 2, value)
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 42
    ws.freeze_panes = "A4"


def ci_system(wb, completed):
    vals = [
        ["Problem", "SMT bridge defects increased from 1.8% to 4.6% during the product ramp-up" if completed else ""],
        ["Scope", "Released kit through inspected board ready for electrical test" if completed else ""],
        ["90-day goal", "Reduce bridge defects to ≤1.5% while keeping cycle time ≤48 s/board" if completed else ""],
        ["Guardrail", "No increase in insufficient-solder defects or first-piece release time" if completed else ""],
    ]
    add_sheet(wb, "Charter", "Improvement Charter", ["Field", "Agreed statement"], vals)
    add_sheet(wb, "Measures", "Measure System", ["Measure", "Baseline", "Target", "Unit", "Owner", "Frequency"], [
        ["Bridge defect rate", .046 if completed else "", .015 if completed else "", "% lots", "Quality Engineer" if completed else "", "Daily" if completed else ""],
        ["First-pass yield", .932 if completed else "", .975 if completed else "", "% boards", "Process Owner" if completed else "", "Daily" if completed else ""],
        ["Cycle time", 51.4 if completed else "", 48 if completed else "", "seconds/board", "Production Lead" if completed else "", "Shift" if completed else ""],
    ])
    add_sheet(wb, "Roles", "Team Role Map", ["Role", "Accountability", "Named owner", "Decision boundary"], [
        ["Sponsor", "Approve priority and remove barriers", "Operations Director" if completed else "", "Funding and escalation" if completed else ""],
        ["Process owner", "Own process performance and sustainment", "SMT Manager" if completed else "", "Standard work and controls" if completed else ""],
        ["Facilitator", "Maintain method and balanced participation", "CI Lead" if completed else "", "Analysis cadence" if completed else ""],
    ])
    add_sheet(wb, "Governance", "Governance Cadence", ["Cadence", "Purpose", "Evidence", "Owner", "Escalation trigger"], [
        ["Daily", "Review defect and stop signals", "Shift board", "Shift Lead", ">3 bridge defects/lot" if completed else ""],
        ["Weekly", "Review actions and learning", "Action register", "CI Lead", "Overdue high-risk action" if completed else ""],
        ["Monthly", "Confirm benefits and controls", "KPI and audit", "Process Owner", "Target missed 2 periods" if completed else ""],
    ])


def sipoc_pdca(wb, completed):
    add_sheet(wb, "Scope", "Process Boundary", ["Field", "Definition", "Rationale"], [
        ["Start", "Released kit, stencil and approved machine program" if completed else "", "Makes preparation controls visible" if completed else ""],
        ["End", "Inspected board released to electrical test" if completed else "", "Captures printing through inspection" if completed else ""],
        ["Out of scope", "PCB design and downstream enclosure assembly" if completed else "", "Separate owners and FMEAs" if completed else ""],
    ])
    add_sheet(wb, "SIPOC", "SIPOC", ["Suppliers", "Inputs", "Process", "Outputs", "Customers", "Requirement"], [
        ["Stores / supplier", "PCB, components, solder paste", "Prepare → print → place → reflow → inspect", "Conforming board", "Electrical test", "FPY ≥97.5%" if completed else ""],
        ["Engineering", "Program, parameters, work instruction", "", "Traceability record", "Quality / customer", "Complete lot history" if completed else ""],
    ])
    add_sheet(wb, "Process Map", "High-Level Process Map", ["Step", "Function", "Output", "Requirement", "Owner"], [
        [1, "Prepare stencil and paste", "Ready printer", "Correct paste and clean aperture" if completed else "", "Operator"],
        [2, "Print solder paste", "Deposited paste", "Volume within specification" if completed else "", "Operator"],
        [3, "Place components", "Populated PCB", "Correct component and alignment" if completed else "", "Operator"],
        [4, "Reflow board", "Soldered PCB", "Profile within validated window" if completed else "", "Process Engineer"],
        [5, "Inspect board", "Dispositioned board", "Detect defined defect classes" if completed else "", "Inspector"],
    ])
    add_sheet(wb, "PDCA", "PDCA Experiment Canvas", ["Stage", "Decision", "Measure", "Owner", "Evidence"], [
        ["Plan", "Trigger stencil cleaning from print-volume trend" if completed else "", "Bridge defects and paste volume" if completed else "", "CI Lead" if completed else "", "Baseline + hypothesis" if completed else ""],
        ["Do", "Pilot on Printer P2 for two weeks" if completed else "", "Representative products/shifts" if completed else "", "Shift Leads" if completed else "", "Pilot log" if completed else ""],
        ["Check", "Compare baseline, outcome and guardrails" if completed else "", "Rate, FPY, cycle time" if completed else "", "Quality Engineer" if completed else "", "Analysis" if completed else ""],
        ["Act", "Scale, revise or stop" if completed else "", "Decision rule" if completed else "", "Process Owner" if completed else "", "Approval" if completed else ""],
    ])


def pareto(wb, completed):
    rng = random.Random(9031)
    defects = ["Solder bridge"] * 42 + ["Insufficient solder"] * 26 + ["Misalignment"] * 15 + ["Wrong component"] * 9 + ["Tombstone"] * 5 + ["Other"] * 3
    rows = []
    for i in range(240):
        category = defects[i % len(defects)] if i < 180 else "None"
        machine = "P2" if category == "Solder bridge" and i % 3 else rng.choice(["P1", "P2"])
        shift = "Late" if category == "Solder bridge" and i % 4 else rng.choice(["Day", "Late"])
        rows.append([f"LOT-{1001+i}", f"2026-07-{1+i%20:02d}", rng.choice(["SC-100", "SC-200"]), machine, shift, category, 1 if category != "None" else 0])
    add_sheet(wb, "Data Dictionary", "Defect Data Dictionary", ["Field", "Definition", "Allowed value / rule"], [
        ["Lot ID", "Unique inspected production lot", "LOT-nnnn"], ["Machine", "Printer used", "P1 / P2"],
        ["Shift", "Production shift", "Day / Late"], ["Defect Category", "Primary inspection result", "Controlled list including None"],
        ["Defect Count", "Count used in summary", "0 or 1"],
    ])
    ws = add_sheet(wb, "Defect Data", "Orion Defect Log", ["Lot ID", "Date", "Product", "Machine", "Shift", "Defect Category", "Defect Count"], rows)
    ws.auto_filter.ref = f"A4:G{ws.max_row}"
    cats = ["Solder bridge", "Insufficient solder", "Misalignment", "Wrong component", "Tombstone", "Other"]
    summary = []
    for idx, cat in enumerate(cats, 5):
        count_formula = f'=COUNTIF(\'Defect Data\'!$F$5:$F$244,A{idx})'
        pct_formula = f'=B{idx}/SUM($B$5:$B$10)'
        cum_formula = f'=SUM($C$5:C{idx})'
        summary.append([cat, count_formula, pct_formula, cum_formula, "Investigate" if completed and idx == 5 else ""])
    ws = add_sheet(wb, "Pareto", "Defect Pareto", ["Defect category", "Count", "% of defects", "Cumulative %", "Decision"], summary)
    for row in range(5, 11):
        ws[f"C{row}"].number_format = "0.0%"; ws[f"D{row}"].number_format = "0.0%"
    bar = BarChart(); bar.title = "Defect Pareto"; bar.y_axis.title = "Count"
    bar.add_data(Reference(ws, min_col=2, min_row=4, max_row=10), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=5, max_row=10))
    line = LineChart(); line.y_axis.axId = 200; line.y_axis.title = "Cumulative %"; line.y_axis.crosses = "max"
    line.add_data(Reference(ws, min_col=4, min_row=4, max_row=10), titles_from_data=True)
    bar += line; ws.add_chart(bar, "G4")
    add_sheet(wb, "Stratification", "Leading-Defect Stratification", ["Dimension", "Group", "Count formula", "Interpretation", "Limitation"], [
        ["Machine", "P2", '=COUNTIFS(\'Defect Data\'!$F$5:$F$244,"Solder bridge",\'Defect Data\'!$D$5:$D$244,"P2")', "Signal concentrated on P2" if completed else "", "Exposure differs by machine" if completed else ""],
        ["Shift", "Late", '=COUNTIFS(\'Defect Data\'!$F$5:$F$244,"Solder bridge",\'Defect Data\'!$E$5:$E$244,"Late")', "Late-shift concentration" if completed else "", "Check product mix" if completed else ""],
    ])


def rca(wb, completed):
    add_sheet(wb, "Fishbone", "6M Fishbone", ["6M branch", "Specific candidate mechanism", "Expected evidence", "Status"], [
        ["People", "Cleaning trigger not understood" if completed else "", "Missed trigger records" if completed else "", "Open" if completed else ""],
        ["Machine", "P2 print-volume sensor drifts" if completed else "", "Gauge check bias" if completed else "", "Rejected" if completed else ""],
        ["Method", "Fixed cleaning interval misses rapid blockage" if completed else "", "Bridges rise before interval" if completed else "", "Confirmed" if completed else ""],
        ["Material", "Paste exposure exceeds limit" if completed else "", "Exposure log correlation" if completed else "", "Open" if completed else ""],
        ["Measurement", "Inspection classification inconsistent" if completed else "", "Attribute agreement" if completed else "", "Rejected" if completed else ""],
        ["Environment", "Late-shift humidity differs" if completed else "", "Humidity trend" if completed else "", "Open" if completed else ""],
    ])
    add_sheet(wb, "5 Whys", "Evidence-Backed 5 Whys", ["Level", "Why?", "Answer / mechanism", "Evidence", "Confidence"], [
        [1, "Why did bridges form?", "Excess paste connected adjacent pads" if completed else "", "AOI images" if completed else "", "High" if completed else ""],
        [2, "Why was excess paste deposited?", "Stencil aperture was partly blocked and release became uneven" if completed else "", "Print-volume trend" if completed else "", "High" if completed else ""],
        [3, "Why did blockage persist?", "Cleaning occurred only at the fixed interval" if completed else "", "Machine log" if completed else "", "High" if completed else ""],
        [4, "Why was the interval not responsive?", "Standard work had no condition-based trigger" if completed else "", "WI review" if completed else "", "High" if completed else ""],
        [5, "Why was no trigger defined?", "PFMEA and control plan were not updated after the new product ramp" if completed else "", "Revision history" if completed else "", "High" if completed else ""],
    ])
    add_sheet(wb, "Fault Tree", "Fault Tree Logic", ["Node ID", "Parent", "Gate", "Event", "Evidence status"], [
        ["T", "", "TOP", "Solder bridge after reflow", "Observed"], ["B1", "T", "OR", "Excess paste deposited", "Supported" if completed else ""],
        ["B2", "T", "OR", "Placement creates conductive path", "Open" if completed else ""], ["B3", "T", "OR", "Paste slumps during reflow", "Open" if completed else ""],
    ])
    add_sheet(wb, "Verification", "Cause Verification Matrix", ["Candidate cause", "Prediction", "Evidence source", "Test", "Decision rule", "Result", "Disposition"], [
        ["Fixed cleaning interval misses blockage", "Paste volume drift precedes bridge defects", "Printer log + AOI", "Compare 20 lots before/after trigger", "≥50% defect reduction without guardrail breach", "63% reduction" if completed else "", "Confirmed" if completed else ""],
    ])
    add_sheet(wb, "Findings", "Verified Findings", ["Finding", "Status", "Implication", "Next action", "Owner"], [
        ["No condition-based cleaning trigger", "Confirmed" if completed else "", "Prevention control gap" if completed else "", "Pilot automated trigger" if completed else "", "Process Engineer" if completed else ""],
    ])


def pfmea(wb, completed):
    add_sheet(wb, "Scope", "PFMEA Scope and Team", ["Field", "Decision / evidence"], [
        ["Purpose", "Prevent SMT solder defects during Orion product ramp" if completed else ""], ["Boundary", "Kit release through inspection release" if completed else ""],
        ["Team", "Process, production, quality, maintenance and operators" if completed else ""], ["Review triggers", "Change, capability signal, incident, complaint or RCA learning" if completed else ""],
    ])
    headers = ["ID", "Process step", "Function / requirement", "Failure mode", "Effect", "S", "Cause", "O", "Prevention control", "Detection control", "D", "RPN", "Action", "Owner", "Due", "New S", "New O", "New D", "New RPN"]
    rows = [
        ["P-01", "Print paste", "Deposit paste within volume specification", "Excess paste", "Bridge causes intermittent signal / customer return", 8 if completed else "", "Stencil aperture partly blocked", 6 if completed else "", "Fixed cleaning interval", "AOI after reflow", 5 if completed else "", "=F5*H5*K5", "Condition-based cleaning trigger" if completed else "", "Process Engineer" if completed else "", "2026-09-15" if completed else "", 8 if completed else "", 3 if completed else "", 4 if completed else "", "=P5*Q5*R5"],
        ["P-02", "Place component", "Place component within alignment tolerance", "Misaligned component", "Open or weak joint at test", 7 if completed else "", "Feeder indexing wear", 4 if completed else "", "Preventive maintenance", "AOI alignment check", 3 if completed else "", "=F6*H6*K6", "Add feeder-life alert" if completed else "", "Maintenance Lead" if completed else "", "2026-09-20" if completed else "", 7 if completed else "", 2 if completed else "", 3 if completed else "", "=P6*Q6*R6"],
        ["P-03", "Reflow", "Heat joints within validated profile", "Incomplete wetting", "Insufficient solder / early-life failure", 8 if completed else "", "Profile drifts after product change", 3 if completed else "", "Recipe approval", "First-piece test", 4 if completed else "", "=F7*H7*K7", "Interlock recipe to product ID" if completed else "", "Controls Engineer" if completed else "", "2026-09-25" if completed else "", 8 if completed else "", 2 if completed else "", 2 if completed else "", "=P7*Q7*R7"],
    ]
    ws = add_sheet(wb, "PFMEA", "Orion Process FMEA", headers, rows, "RPN = Severity × Occurrence × Detection; higher D means weaker detection")
    for col in range(1, 20): ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions["C"].width = 28; ws.column_dimensions["E"].width = 28; ws.column_dimensions["G"].width = 28
    ws.conditional_formatting.add("L5:L50", ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="F8696B"))
    guides = []
    for score in range(1, 11):
        guides.append([score, "Minor / local" if score <= 3 else "Moderate" if score <= 6 else "Major / customer / safety", "Rare" if score <= 3 else "Occasional" if score <= 6 else "Frequent", "Almost certain detection" if score <= 3 else "Moderate detection" if score <= 6 else "Weak / unlikely detection"])
    add_sheet(wb, "Rating Guides", "Internal S/O/D Learning Guide", ["Score", "Severity guide", "Occurrence guide", "Detection guide"], guides, "Illustrative course guide only; use the organisation's approved scale in real work")
    add_sheet(wb, "Risk Review", "Severity-First Risk Review", ["PFMEA ID", "S", "O", "D", "RPN", "Why action is / is not required"], [
        ["P-01", "='PFMEA'!F5", "='PFMEA'!H5", "='PFMEA'!K5", "='PFMEA'!L5", "Customer return and weak prevention control" if completed else ""],
    ])
    add_sheet(wb, "Actions", "Action Register", ["PFMEA ID", "Action", "Hierarchy", "Owner", "Due", "Implementation evidence", "Effectiveness evidence", "Status"], [
        ["P-01", "Condition-based stencil-cleaning trigger" if completed else "", "Prevention" if completed else "", "Process Engineer" if completed else "", "2026-09-15" if completed else "", "Approved program + WI" if completed else "", "Pilot comparison" if completed else "", "Planned" if completed else ""],
    ])
    add_sheet(wb, "Re-score", "Post-Action Re-score", ["PFMEA ID", "Old RPN", "Verified change", "New S", "New O", "New D", "New RPN", "Residual decision"], [
        ["P-01", "='PFMEA'!L5", "Trigger implemented and pilot effective" if completed else "", "='PFMEA'!P5", "='PFMEA'!Q5", "='PFMEA'!R5", "='PFMEA'!S5", "Monitor in Control Plan" if completed else ""],
    ])


def selection(wb, completed):
    options = ["Condition-based cleaning", "First-piece verification", "New stencil", "100% manual inspection", "Operator refresher"]
    add_sheet(wb, "Options", "Improvement Options", ["Option", "PFMEA risk", "Mechanism targeted", "Evidence / assumption"], [[o, "P-01" if i < 4 else "P-02", "Aperture blockage / escape" if completed else "", "RCA and supplier estimate" if completed else ""] for i, o in enumerate(options)])
    criteria = [["Risk reduction", 30], ["Effectiveness", 20], ["Feasibility", 15], ["Speed", 10], ["Cost", 10], ["Side effects", 5], ["Sustainability", 10]]
    add_sheet(wb, "Criteria", "Weighted Criteria", ["Criterion", "Weight %"], criteria)
    rows = []
    for idx, option in enumerate(options, 5):
        scores = [5, 5, 4, 4, 3, 4, 5] if idx == 5 else [4, 4, 4, 5, 4, 4, 4]
        row = [option] + (scores if completed else ["" for _ in scores])
        row += [f"=SUM(B{idx}*'Criteria'!$B$5,C{idx}*'Criteria'!$B$6,D{idx}*'Criteria'!$B$7,E{idx}*'Criteria'!$B$8,F{idx}*'Criteria'!$B$9,G{idx}*'Criteria'!$B$10,H{idx}*'Criteria'!$B$11)/100"]
        rows.append(row)
    ws = add_sheet(wb, "Weighted Matrix", "Weighted Decision Matrix", ["Option", "Risk reduction", "Effectiveness", "Feasibility", "Speed", "Cost", "Side effects", "Sustainability", "Weighted total"], rows)
    ws.column_dimensions["A"].width = 28
    add_sheet(wb, "Pugh", "Pugh Matrix", ["Criterion", "Datum: current method", *options], [[c[0], "S", *("+" if completed and i < 2 else "S" for i in range(5))] for c in criteria])
    add_sheet(wb, "Cost Benefit", "Cost–Benefit Ranges", ["Option", "Low cost", "High cost", "Low annual benefit", "High annual benefit", "Key uncertainty"], [[o, 4000 + i*1500, 7000 + i*2200, 12000 + i*3500, 25000 + i*5000, "Adoption and defect exposure" if completed else ""] for i, o in enumerate(options)])
    add_sheet(wb, "Recommendation", "Recommendation", ["Rank", "Option", "Decision", "Rationale", "Dependency / risk"], [[1, "Condition-based cleaning" if completed else "", "Pilot" if completed else "", "Best prevention and risk-reduction balance" if completed else "", "Validate sensor and operator response" if completed else ""], [2, "First-piece verification" if completed else "", "Pilot" if completed else "", "Strengthens early detection while prevention matures" if completed else "", "Avoid release delay" if completed else ""]])


def pilot(wb, completed):
    add_sheet(wb, "Charter", "Pilot Charter", ["Field", "Definition", "Owner"], [["Hypothesis", "A print-volume trigger prompts cleaning before aperture blockage creates bridges" if completed else "", "Process Engineer" if completed else ""], ["Scope", "Printer P2, two products, both shifts, two weeks" if completed else "", "SMT Manager" if completed else ""], ["Decision", "Scale if ≥50% bridge reduction with no guardrail breach" if completed else "", "Process Owner" if completed else ""]])
    add_sheet(wb, "Measures", "Pilot Measures", ["Type", "Measure", "Baseline", "Target", "Definition", "Frequency"], [["Outcome", "Bridge defect rate", .046 if completed else "", .023 if completed else "", "Bridge lots / inspected lots" if completed else "", "Daily"], ["Leading", "Cleaning trigger compliance", "" if not completed else .0, .95 if completed else "", "Completed triggers / generated triggers" if completed else "", "Shift"], ["Guardrail", "Cycle time", 51.4 if completed else "", 52 if completed else "", "Average seconds/board" if completed else "", "Daily"]])
    add_sheet(wb, "Design", "Pilot Design and Reaction", ["Element", "Decision", "Reason / safeguard"], [["Products", "SC-100 and SC-200" if completed else "", "Representative aperture patterns" if completed else ""], ["Stop rule", "Stop if defects exceed 6% or safety interlock is bypassed" if completed else "", "Contain risk" if completed else ""], ["Rollback", "Restore validated fixed interval and contain suspect lots" if completed else "", "Reversible change" if completed else ""]])
    add_sheet(wb, "Change Control", "Temporary Change Control", ["Change", "Approval", "Document", "Training", "Status"], [["Condition-based cleaning trigger" if completed else "", "Engineering + Quality" if completed else "", "Temporary WI-17" if completed else "", "Operators and shift leads" if completed else "", "Approved" if completed else ""]])
    add_sheet(wb, "Coaching", "Coaching Observation", ["Stage", "Coach behaviour", "Learner evidence", "Result / feedback"], [["Explain", "Connect trigger to defect mechanism" if completed else "", "Operator explains why" if completed else "", ""], ["Demonstrate", "Show trigger response and record" if completed else "", "Operator observes full cycle" if completed else "", ""], ["Practice", "Operator responds to simulated trigger" if completed else "", "Correct action without prompt" if completed else "", ""], ["Observe", "Watch normal production response" if completed else "", "Adherence record" if completed else "", ""], ["Feedback", "Reinforce and correct specifically" if completed else "", "Operator confirms next action" if completed else "", ""]])
    add_sheet(wb, "Decision", "Scale / Revise / Stop", ["Outcome", "Guardrails", "Adoption", "Decision", "Rationale", "Approval"], [["63% bridge reduction" if completed else "", "No cycle-time breach" if completed else "", "96% trigger compliance" if completed else "", "Scale" if completed else "", "Meets rule across both shifts" if completed else "", "Process Owner" if completed else ""]])


def spc(wb, completed):
    values = [99.7,100.5,100.1,99.8,100.4,100.2,99.9,100.3,100.1,99.6,100.4,100.6,100.2,100.0,99.7,100.3,100.5,100.1,99.9,100.2,100.4,100.7,101.5,100.6,100.3,100.2,100.1,99.8,100.0,100.3]
    add_sheet(wb, "Data Dictionary", "SPC Data Dictionary", ["Field", "Definition", "Rule"], [["Day", "Time order of daily subgroup mean", "1–30"], ["Paste volume index", "Normalised daily mean", "Specification 98.5–101.5"], ["Context", "Shift, material or maintenance note", "Record before interpretation"]])
    rows = []
    for i, value in enumerate(values, 5):
        rows.append([i-4, value, "=AVERAGE($B$5:$B$34)", "=STDEV($B$5:$B$34)", "=C5+3*D5", "=MAX(0,C5-3*D5)", f'=IF(OR(B{i}>E{i},B{i}<F{i}),"Beyond limit","")', "Maintenance adjustment" if completed and i == 27 else ""])
    ws = add_sheet(wb, "Control Chart", "Paste-Volume Control Chart", ["Day", "Value", "Centre line", "Std dev", "UCL", "LCL", "Signal", "Context"], rows)
    line = LineChart(); line.title = "Time-Ordered Process Behaviour"; line.y_axis.title = "Paste volume index"; line.x_axis.title = "Day"
    line.add_data(Reference(ws, min_col=2, max_col=6, min_row=4, max_row=34), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=5, max_row=34)); ws.add_chart(line, "J4")
    add_sheet(wb, "Signal Log", "Signal Investigation Log", ["Date / day", "Signal", "Containment", "Cause evidence", "Disposition", "PFMEA update"], [[23, "Beyond UCL" if completed else "", "Hold affected lot" if completed else "", "Maintenance adjustment shifted setting" if completed else "", "Restore and verify" if completed else "", "Review P-01 occurrence/control" if completed else ""]])
    add_sheet(wb, "Capability", "Capability Check", ["Measure", "Formula / value", "Interpretation"], [["USL", 101.5, "Upper specification"], ["LSL", 98.5, "Lower specification"], ["Mean", "=AVERAGE('Control Chart'!B5:B34)", "Use after stability review"], ["Std dev", "=STDEV('Control Chart'!B5:B34)", "Within the observed period"], ["Cp", "=(B5-B6)/(6*B8)", "Potential capability"], ["Cpk", "=MIN((B5-B7)/(3*B8),(B7-B6)/(3*B8))", "Accounts for centring"]])
    add_sheet(wb, "Evaluation", "Improvement Evaluation", ["Dimension", "Finding", "Evidence", "Next action"], [["Effectiveness", "Bridge defects reduced 63%" if completed else "", "Pilot comparison" if completed else "", "Scale with control" if completed else ""], ["Stability", "One special-cause signal investigated" if completed else "", "Control chart + log" if completed else "", "Verify after maintenance" if completed else ""], ["Capability", "Interpret only after stable period" if completed else "", "Cp/Cpk sheet" if completed else "", "Collect 20 more days" if completed else ""]])


def control_plan(wb, completed):
    add_sheet(wb, "Traceability", "PFMEA-to-Control Traceability", ["PFMEA ID", "Risk / characteristic", "Control Plan row", "Control type", "Gap"], [["P-01", "Paste volume / bridge risk", "CP-01", "Prevention + detection" if completed else "", "None" if completed else ""], ["P-02", "Placement alignment", "CP-02", "Prevention + detection" if completed else "", "Feeder-life alert pending" if completed else ""]])
    add_sheet(wb, "Control Plan", "Orion Process Control Plan", ["ID", "Process step", "Characteristic", "Specification / target", "Method", "Sample / frequency", "Owner", "Record", "Reaction plan ID", "PFMEA ID"], [["CP-01", "Print paste", "Paste volume index", "98.5–101.5", "Printer sensor + trend", "Every board / shift review", "Operator / Shift Lead", "Printer log", "RP-01", "P-01"], ["CP-02", "Place component", "Alignment", "Within AOI tolerance", "AOI", "100%", "Operator", "AOI log", "RP-02", "P-02"]])
    add_sheet(wb, "Reaction Plan", "Reaction Plan", ["ID", "Trigger", "Stop / contain", "Suspect boundary", "Notify", "Correct / verify", "Release authority"], [["RP-01", "Paste signal or bridge >3/lot", "Stop printer; hold lot", "Since last accepted verification", "Shift, Quality, Process", "Clean/restore; first-piece verification", "Quality Engineer"], ["RP-02", "AOI alignment failure", "Stop placement; hold lot", "Since last accepted AOI", "Shift, Quality, Maintenance", "Check feeder; verify 5 boards", "Quality Engineer"]])
    add_sheet(wb, "Readiness", "Control Readiness Checklist", ["Check", "Evidence", "Owner", "Ready?", "Gap / action"], [["Measurement validated", "Sensor MSA record" if completed else "", "Quality" if completed else "", "Yes" if completed else "", ""], ["All shifts trained", "Attendance + observation" if completed else "", "Production" if completed else "", "Yes" if completed else "", ""], ["Escalation tested", "Reaction simulation" if completed else "", "Shift Lead" if completed else "", "Yes" if completed else "", ""]])
    add_sheet(wb, "Review Calendar", "Control and PFMEA Review Calendar", ["Cadence / trigger", "Review focus", "Evidence", "Owner", "Action if triggered"], [["Monthly", "Trend, audit and action status", "Control charts + audit", "Process Owner", "Update PFMEA/control plan"], ["Engineering change", "New failure paths", "Change record", "Engineering", "Cross-functional review"], ["Capability loss / incident / complaint", "Contain, RCA and residual risk", "Signal or case record", "Quality", "Immediate PFMEA review"]])


BUILDERS = {1: ci_system, 2: sipoc_pdca, 3: pareto, 4: rca, 5: pfmea, 6: pfmea, 7: selection, 8: pilot, 9: spc, 10: control_plan}


def build_workbook(activity, completed, path):
    wb = Workbook()
    instructions_sheet(wb, activity)
    BUILDERS[activity["num"]](wb, completed)
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_margins.left = .25; ws.page_margins.right = .25
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def activity_markdown(activity):
    lines = [
        f"# Activity {activity['num']:02d} — {activity['title']}", "",
        f"**Course:** {C.TITLE} ({C.COURSE_CODE})  ", f"**Assessment mapping:** {activity['codes']}  ",
        f"**Objective:** {activity['objective']}", "", "## Scenario", "", activity["scenario"], "",
        "## Deliverable", "", activity["build"], "", "## Files", "",
        "- `Starter.xlsx` — your working file.", "- `Completed.xlsx` — a worked comparison; review after attempting the activity.",
        "- `Activity-Brief.pdf` — printable scenario, questions and submission check.", "- `Checklist.md` — final verification.",
        "", "## Detailed procedure", "",
    ]
    for i, (step, evidence) in enumerate(activity["steps"], 1):
        lines += [f"### Step {i} — {step}", "", f"Open or update: **{evidence}**.", "",
                  "Record your rationale, source evidence and any assumption directly in the workbook. Keep the entry specific enough for another team member to reproduce your decision.", ""]
    lines += ["## Scenario questions", ""]
    lines += [f"{i}. {question}" for i, question in enumerate(activity["questions"], 1)]
    lines += ["", "## Acceptance criterion", "", activity["test"], "", "## Submission", "",
              "Save the completed workbook with your name or learner ID in the filename. Submit the workbook and any requested screenshots or written responses through the approved course channel.", ""]
    if activity["num"] == 2:
        lines += ["Optional browser tool: https://alfredang.github.io/sipoc/", ""]
    if activity["num"] == 3:
        lines += ["Optional browser tool: https://alfredang.github.io/paretochart/", ""]
    if activity["num"] == 4:
        lines += ["Optional browser tools: https://alfredang.github.io/fishbone/ and https://alfredang.github.io/5whys/", ""]
    if activity["num"] == 10:
        lines += ["Optional browser tool for feedback relationships: https://alfredang.github.io/systemloop/", ""]
    return "\n".join(lines)


def checklist(activity):
    items = [
        "I used the Orion scenario and did not invent an unsupported cause.",
        "I completed every required workbook field and retained the formulas.",
        "My evidence, assumptions and decision rationale are traceable.",
        "I answered both scenario questions.",
        "I applied the activity acceptance criterion.",
        "I saved the completed file with a clear filename.",
    ]
    return "\n".join([f"# Activity {activity['num']:02d} Checklist", "", f"**Acceptance criterion:** {activity['test']}", ""] + [f"- [ ] {item}" for item in items] + [""])


def build_pdf(activity, path):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CourseTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=27, textColor=colors.HexColor("#161B26"), alignment=TA_CENTER, spaceAfter=14))
    styles.add(ParagraphStyle(name="BlueHead", parent=styles["Heading2"], fontName="Helvetica-Bold", textColor=colors.HexColor("#1F6FEB"), spaceBefore=8, spaceAfter=5))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=16*mm, bottomMargin=16*mm,
                            title=f"Activity {activity['num']:02d} - {activity['title']}")
    story = [Paragraph("TERTIARY INFOTECH ACADEMY PTE LTD", styles["Heading3"]), Spacer(1, 8),
             Paragraph(f"Activity {activity['num']:02d}", styles["Heading2"]), Paragraph(activity["title"], styles["CourseTitle"]),
             Paragraph(f"{C.TITLE} · {C.COURSE_CODE}", styles["Normal"]), Spacer(1, 12),
             Paragraph("Scenario", styles["BlueHead"]), Paragraph(activity["scenario"], styles["BodyText"]),
             Paragraph("Objective", styles["BlueHead"]), Paragraph(activity["objective"], styles["BodyText"]),
             Paragraph("Deliverable", styles["BlueHead"]), Paragraph(activity["build"], styles["BodyText"]),
             PageBreak(), Paragraph("Scenario questions", styles["BlueHead"])]
    for i, question in enumerate(activity["questions"], 1):
        story += [Paragraph(f"{i}. {question}", styles["BodyText"]), Spacer(1, 26)]
    story += [Paragraph("Submission checklist", styles["BlueHead"])]
    table_data = [["✓", "Check"], *[["☐", item] for item in ["Workbook complete", "Formulas retained", "Evidence and assumptions recorded", "Questions answered", "Acceptance criterion applied"]]]
    table = Table(table_data, colWidths=[14*mm, 145*mm])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1F6FEB")), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
                               ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#D9E2F0")),
                               ("VALIGN", (0,0), (-1,-1), "TOP"), ("FONTNAME", (0,1), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), 9),
                               ("BOTTOMPADDING", (0,0), (-1,-1), 7), ("TOPPADDING", (0,0), (-1,-1), 7)]))
    story += [table, Spacer(1, 14), Paragraph(f"Acceptance criterion: {activity['test']}", styles["BodyText"])]
    doc.build(story)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    index = [f"# Activities — {C.TITLE}", "", "Ten connected activities follow Orion Precision Electronics from problem definition through a sustained Control Plan.", "",
             "Detailed procedures appear in each folder and in the Learner Guide.", ""]
    for activity in ACTIVITIES:
        folder = OUT / f"Activity {activity['num']:02d} - {safe_title(activity['title'])}"
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "README.md").write_text(activity_markdown(activity), encoding="utf-8")
        (folder / "Checklist.md").write_text(checklist(activity), encoding="utf-8")
        build_pdf(activity, folder / "Activity-Brief.pdf")
        build_workbook(activity, False, folder / "Starter.xlsx")
        build_workbook(activity, True, folder / "Completed.xlsx")
        index.append(f"{activity['num']}. [{activity['title']}]({folder.name}/README.md) — {activity['codes']}")
    (OUT / "README.md").write_text("\n".join(index) + "\n", encoding="utf-8")
    print(f"Built {len(ACTIVITIES)} activity folders in {OUT}")


if __name__ == "__main__":
    main()
