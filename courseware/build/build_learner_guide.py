#!/usr/bin/env python3
"""Build the aligned Markdown and DOCX Learner Guide."""
import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import course_data as C
import prodoc
from data_domain1 import DOMAIN1
from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5

ACTIVITIES = DOMAIN1 + DOMAIN2 + DOMAIN3 + DOMAIN4 + DOMAIN5
ROOT = HERE.parents[1]
OUT = ROOT / "courseware"
LOGO = OUT / "assets" / "tertiary-infotech-logo.png"
BLUE = RGBColor(0x1F, 0x6F, 0xEB)
TEAL = RGBColor(0x10, 0xB9, 0x81)
GREY = RGBColor(0x5B, 0x63, 0x72)

SOURCE_NOTES = {
    "FMEA is prevention made visible": "ASQ defines FMEA as a systematic, proactive method for evaluating where and how a process might fail and the relative impact of those failures.",
    "Seven-step FMEA method": "This guide uses the seven-step planning, structure, function, failure, risk, optimisation and documentation sequence described in current AIAG/VDA-style practice.",
    "RPN is useful—but incomplete": "Severity, occurrence and detection are ordinal judgements. Similar RPN values can represent very different risk profiles, so severity and action priority remain visible.",
    "Prospective FMEA and retrospective RCA": "FMEA anticipates credible future failure; RCA explains evidence from an observed event. Verified RCA learning must update the FMEA.",
}


def plain(text):
    return str(text).replace("≤", "less than or equal to ").replace("≥", "greater than or equal to ")


def concept_body(spec):
    kind, title, payload = spec[0], spec[1], spec[2]
    if kind == "big":
        return [payload]
    if kind == "two":
        return ["First perspective: " + "; ".join(payload) + ".", "Contrast or design check: " + "; ".join(spec[3]) + "."]
    if kind == "cards":
        return [f"{label}: " + "; ".join(items) + "." for label, items in payload]
    if kind in {"tiles"}:
        return [f"{label}: {body}" for label, body in payload]
    if kind in {"flow", "pdca", "failurechain", "riskmatrix", "controlplan", "faulttree", "fishbone"}:
        return [" → ".join(payload)]
    if kind == "pareto":
        labels = spec[3]
        total = sum(payload)
        return [f"Illustrative defect mix: {', '.join(f'{label} {value}' for label, value in zip(labels, payload))}. The categories sum to {total}; use the cumulative percentage to focus investigation without claiming the largest category is a root cause."]
    if kind == "controlchart":
        return ["Plot observations in time order, calculate a centre line and statistically derived control limits, then investigate non-random signals. Control limits describe observed behaviour; specification limits describe requirements."]
    return [str(payload)]


def build_markdown():
    lines = [
        f"# Learner Guide — {C.TITLE}", "", f"**Course code:** {C.COURSE_CODE}  ",
        f"**TSC:** {C.TSC_TITLE} ({C.TSC_CODE})  ", f"**Version:** {C.VERSION} — {C.VERSION_DATE}  ",
        "**Duration:** 2 days / 16 hours, including 60-minute Written Assessment and 90-minute Practical Performance", "",
        "## How to Use This Guide", "",
        "Use the concept chapters before each activity. Complete the Starter workbook in the matching activity folder, record evidence and assumptions, then use the checklist. The completed workbook is for comparison after your own attempt.", "",
        "## Connected Case Study", "", f"**{C.CASE_STUDY['company']}** manufactures industrial sensor-controller boards. Its process is {C.CASE_STUDY['process'].lower()}.", "",
        C.CASE_STUDY["problem"] + ".", "", "The improvement goal is to " + C.CASE_STUDY["goal"].lower().replace("pfmea", "PFMEA").replace("control plan", "Control Plan") + ". Every activity advances this same case, so later decisions must trace to evidence created earlier.", "",
        "## Learning Outcomes and Competency Mapping", "",
    ]
    lines += [f"- {item}" for item in C.LEARNING_OUTCOMES]
    lines += ["", "### Knowledge", ""] + [f"- {item}" for item in C.KNOWLEDGE]
    lines += ["", "### Abilities", ""] + [f"- {item}" for item in C.ABILITIES]
    for topic in C.TOPICS:
        lines += ["", f"# Topic {topic['num']} — {topic['title']}", "", f"**Alignment:** {topic['weighting']}  ", f"**Focus:** {topic['subtitle']}", "",
                  "## Topic Orientation", ""]
        for label, body in topic["concepts"]:
            lines += [f"### {label}", "", body, ""]
        lines += ["## Concept Guide", ""]
        for spec in C.CONTENT_SLIDES[topic["num"]]:
            title = spec[1]
            lines += [f"### {title}", ""]
            for paragraph in concept_body(spec):
                lines += [paragraph, ""]
            if title in SOURCE_NOTES:
                lines += [f"**Source note:** {SOURCE_NOTES[title]}", ""]
        for activity in [a for a in ACTIVITIES if a["topic"] == topic["num"]]:
            lines += [f"## Activity {activity['num']:02d} — {activity['title']}", "",
                      f"**Mapping:** {activity['codes']}  ", f"**Objective:** {activity['objective']}", "",
                      "### Scenario", "", activity["scenario"], "", "### Deliverable", "", activity["build"], "",
                      "### Files and tools", "", activity["services"], "", "### Detailed step-by-step procedure", ""]
            for idx, (step, evidence) in enumerate(activity["steps"], 1):
                lines += [f"#### Step {idx} — {step}", "", f"Work in **{evidence}**.", "",
                          "Record the supporting evidence, assumption and decision rationale in the indicated file. If the evidence is incomplete, label the item as an open question and assign the next verification action rather than presenting it as confirmed.", ""]
            lines += ["### Scenario questions", ""] + [f"{i}. {q}" for i, q in enumerate(activity["questions"], 1)]
            lines += ["", "### Acceptance criterion", "", activity["test"], "", "### Submission check", "",
                      "- [ ] All required tabs or sections are complete.", "- [ ] Formula results reconcile to their source data.",
                      "- [ ] Evidence and assumptions are traceable.", "- [ ] Both scenario questions are answered.",
                      "- [ ] The workbook filename clearly identifies the learner.", ""]
            if activity["num"] == 2:
                lines += ["Optional tool: [SIPOC Builder](https://alfredang.github.io/sipoc/)", ""]
            if activity["num"] == 3:
                lines += ["Optional tool: [Pareto Chart Builder](https://alfredang.github.io/paretochart/)", ""]
            if activity["num"] == 4:
                lines += ["Optional tools: [Fishbone Builder](https://alfredang.github.io/fishbone/) and [5 Whys Builder](https://alfredang.github.io/5whys/)", ""]
            if activity["num"] == 10:
                lines += ["Optional tool: [System Loop Builder](https://alfredang.github.io/systemloop/)", ""]
    lines += ["", "# Assessment Readiness", "",
              "The Written Assessment contains four open-ended questions mapped to K1–K4 and takes 60 minutes. The Practical Performance contains five connected tasks mapped to A1–A10 and takes 90 minutes.", "",
              "Explain the reasoning behind your decisions. A correct-looking template without a coherent failure chain, reliable formulas, evidence, ownership and control is incomplete.", "",
              "# References", "",
              "- ASQ. [Failure Mode and Effects Analysis](https://asq.org/quality-resources/fmea).",
              "- Jama Software. [FMEA and regulatory/industry practice](https://www.jamasoftware.com/requirements-management-guide/meeting-regulatory-compliance-and-industry-standards/fmea/).",
              "- Quality-One. [Failure Mode and Effects Analysis](https://quality-one.com/fmea/).",
              "- Rah et al. [Overview of Failure Mode and Effects Analysis: A Patient Safety Tool](https://pmc.ncbi.nlm.nih.gov/articles/PMC10229026/).",
              "- IPH Hannover. [FMEA in manufacturing](https://www.iph-hannover.de/de/dienstleistungen/fertigungsverfahren/fmea/).",
              "- ScienceDirect Topics. [Failure Modes and Effects Analysis](https://www.sciencedirect.com/topics/engineering/failure-modes-and-effects-analysis).",
              "- Wikipedia. [Failure mode and effects analysis](https://en.wikipedia.org/wiki/Failure_mode_and_effects_analysis) and [FMEA](https://de.wikipedia.org/wiki/FMEA).",
              ""]
    return "\n".join(lines)


def set_doc_defaults(doc):
    style = doc.styles["Normal"]
    style.font.name = "Arial"; style.font.size = Pt(10.5); style.font.color.rgb = RGBColor(0x16, 0x1B, 0x26)
    style.paragraph_format.space_after = Pt(5)
    for section in doc.sections:
        section.top_margin = Inches(.72); section.bottom_margin = Inches(.75)
        section.left_margin = Inches(.8); section.right_margin = Inches(.8)
    prodoc.style_headings(doc)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_link(paragraph, text, url):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1F6FEB")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    props.extend([color, underline]); run.append(props)
    node = OxmlElement("w:t"); node.text = text; run.append(node); hyperlink.append(run); paragraph._p.append(hyperlink)


def activity_table(doc, activity):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in [("Mapping", activity["codes"]), ("Objective", activity["objective"]), ("Deliverable", activity["build"]), ("Tools/files", activity["services"])]:
        cells = table.add_row().cells
        cells[0].text = label; cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def build_docx():
    doc = Document()
    set_doc_defaults(doc)
    prodoc.add_cover_page(doc, "Learner Guide", C.TITLE, C.VERSION, org_logo=str(LOGO), course_code=C.COURSE_CODE)
    prodoc.add_version_control(doc, [("v10.0", "Historical reference", "Legacy slide-based guide", C.ORG),
                                     (C.VERSION, C.VERSION_DATE, "Rebuilt as source-aligned concept guide with ten detailed Orion activities", C.ORG)])
    prodoc.add_toc(doc, "1-3")
    doc.add_heading("How to Use This Guide", level=1)
    doc.add_paragraph("Use the concept chapters before each activity. Complete the Starter workbook in the matching activity folder, record evidence and assumptions, then use the checklist. The Completed workbook is for comparison after your own attempt.")
    doc.add_heading("Connected Case Study", level=1)
    p = doc.add_paragraph(); p.add_run(C.CASE_STUDY["company"] + ". ").bold = True; p.add_run(C.CASE_STUDY["process"] + ".")
    doc.add_paragraph(C.CASE_STUDY["problem"] + ".")
    doc.add_paragraph("The goal is to " + C.CASE_STUDY["goal"].lower().replace("pfmea", "PFMEA").replace("control plan", "Control Plan") + ". Every activity advances the same case, so later decisions must trace to earlier evidence.")
    doc.add_heading("Learning Outcomes and Competency Mapping", level=1)
    for item in C.LEARNING_OUTCOMES: add_bullet(doc, item)
    doc.add_heading("Knowledge", level=2)
    for item in C.KNOWLEDGE: add_bullet(doc, item)
    doc.add_heading("Abilities", level=2)
    for item in C.ABILITIES: add_bullet(doc, item)
    for topic in C.TOPICS:
        doc.add_page_break()
        doc.add_heading(f"Topic {topic['num']} — {topic['title']}", level=1)
        p = doc.add_paragraph(); p.add_run("Alignment: ").bold = True; p.add_run(topic["weighting"])
        p = doc.add_paragraph(); p.add_run("Focus: ").bold = True; p.add_run(topic["subtitle"])
        doc.add_heading("Topic Orientation", level=2)
        for label, body in topic["concepts"]:
            doc.add_heading(label, level=3); doc.add_paragraph(body)
        doc.add_heading("Concept Guide", level=2)
        for spec in C.CONTENT_SLIDES[topic["num"]]:
            title = spec[1]
            doc.add_heading(title, level=3)
            for paragraph in concept_body(spec): doc.add_paragraph(paragraph)
            if title in SOURCE_NOTES:
                p = doc.add_paragraph(); r = p.add_run("Source note: "); r.bold = True; r.font.color.rgb = TEAL; p.add_run(SOURCE_NOTES[title])
        for activity in [a for a in ACTIVITIES if a["topic"] == topic["num"]]:
            doc.add_page_break()
            doc.add_heading(f"Activity {activity['num']:02d} — {activity['title']}", level=2)
            activity_table(doc, activity)
            doc.add_heading("Scenario", level=3); doc.add_paragraph(activity["scenario"])
            doc.add_heading("Detailed Step-by-Step Procedure", level=3)
            for idx, (step, evidence) in enumerate(activity["steps"], 1):
                p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(3); r = p.add_run(f"Step {idx} — {step}"); r.bold = True; r.font.color.rgb = BLUE
                p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(3); p.add_run("Work in: ").bold = True; p.add_run(evidence)
                p = doc.add_paragraph("Record the supporting evidence, assumption and decision rationale in the indicated file. If evidence is incomplete, label it as open and assign the next verification action rather than presenting it as confirmed.")
                p.paragraph_format.space_after = Pt(4)
            doc.add_heading("Scenario Questions", level=3)
            for q in activity["questions"]: add_number(doc, q)
            doc.add_heading("Acceptance Criterion", level=3); doc.add_paragraph(activity["test"])
            doc.add_heading("Submission Check", level=3)
            for text in ["All required tabs or sections are complete.", "Formula results reconcile to their source data.", "Evidence and assumptions are traceable.", "Both scenario questions are answered.", "The workbook filename clearly identifies the learner."]:
                add_bullet(doc, "☐ " + text)
            urls = {2: ("SIPOC Builder", "https://alfredang.github.io/sipoc/"), 3: ("Pareto Chart Builder", "https://alfredang.github.io/paretochart/"), 4: ("Fishbone Builder", "https://alfredang.github.io/fishbone/"), 10: ("System Loop Builder", "https://alfredang.github.io/systemloop/")}
            if activity["num"] in urls:
                p = doc.add_paragraph("Optional tool: "); add_link(p, *urls[activity["num"]])
            if activity["num"] == 4:
                p = doc.add_paragraph("Optional companion tool: "); add_link(p, "5 Whys Builder", "https://alfredang.github.io/5whys/")
    doc.add_page_break(); doc.add_heading("Assessment Readiness", level=1)
    doc.add_paragraph("The Written Assessment contains four open-ended questions mapped to K1–K4 and takes 60 minutes. The Practical Performance contains five connected tasks mapped to A1–A10 and takes 90 minutes.")
    doc.add_paragraph("Explain the reasoning behind your decisions. A correct-looking template without a coherent failure chain, reliable formulas, evidence, ownership and control is incomplete.")
    doc.add_heading("References", level=1)
    sources = [("ASQ — Failure Mode and Effects Analysis", "https://asq.org/quality-resources/fmea"),
               ("Jama Software — FMEA and regulatory/industry practice", "https://www.jamasoftware.com/requirements-management-guide/meeting-regulatory-compliance-and-industry-standards/fmea/"),
               ("Quality-One — Failure Mode and Effects Analysis", "https://quality-one.com/fmea/"),
               ("PMC — Overview of FMEA as a patient safety tool", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10229026/"),
               ("IPH Hannover — FMEA in manufacturing", "https://www.iph-hannover.de/de/dienstleistungen/fertigungsverfahren/fmea/"),
               ("ScienceDirect Topics — Failure Modes and Effects Analysis", "https://www.sciencedirect.com/topics/engineering/failure-modes-and-effects-analysis"),
               ("Wikipedia — Failure mode and effects analysis", "https://en.wikipedia.org/wiki/Failure_mode_and_effects_analysis")]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet"); add_link(p, label, url)
    prodoc.add_page_numbers(doc, f"{C.TITLE} · {C.COURSE_CODE}")
    prodoc.enable_update_fields(doc)
    path = OUT / f"LG-{C.SHORT_TITLE}-{C.VERSION}.docx"
    doc.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    md_path = OUT / f"LG-{C.SHORT_TITLE}.md"
    md_path.write_text(build_markdown(), encoding="utf-8")
    docx_path = build_docx()
    print(f"Saved {md_path}")
    print(f"Saved {docx_path}")


if __name__ == "__main__":
    main()
