#!/usr/bin/env python3
"""Build the visual FMEA course deck from the shared content modules.

The deck explains concepts and decisions. Detailed activity procedures live in
the Learner Guide and the individual activity folders.
"""
import json
import math
import os
import re
import sys
import tempfile
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
from data_domain1 import DOMAIN1
from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5

ACTIVITIES = DOMAIN1 + DOMAIN2 + DOMAIN3 + DOMAIN4 + DOMAIN5
REPO = os.path.dirname(os.path.dirname(HERE))
ASSETS = os.path.join(REPO, "courseware", "assets")

BLUE = RGBColor(0x1F, 0x6F, 0xEB)
TEAL = RGBColor(0x10, 0xB9, 0x81)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
RED = RGBColor(0xDC, 0x26, 0x26)
DEEP = RGBColor(0x0F, 0x3D, 0x3E)
INK = RGBColor(0x16, 0x1B, 0x26)
GREY = RGBColor(0x5B, 0x63, 0x72)
LIGHT = RGBColor(0xF5, 0xF8, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
PALETTE = [BLUE, TEAL, VIOLET, AMBER]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
PAGE = {"n": 0}


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None):
    shape = s.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def oval(s, x, y, w, h, color):
    shape = s.shapes.add_shape(9, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=4):
    box = s.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for i, line in enumerate(runs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        for text, size, color, bold in line:
            run = p.add_run()
            run.text = text
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def bullets(s, x, y, w, h, items, size=17, gap=9):
    box = s.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = "•  " + str(item)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = INK
    return box


def footer(s):
    PAGE["n"] += 1
    txt(s, Inches(.42), Inches(7.05), Inches(7), Inches(.25),
        [[(f"{C.SHORT_TITLE}  ·  {C.COURSE_CODE}", 9, GREY, False)]])
    txt(s, Inches(5.0), Inches(7.05), Inches(3.4), Inches(.25),
        [[("© 2026 Tertiary Infotech Academy Pte Ltd", 9, GREY, False)]], PP_ALIGN.CENTER)
    txt(s, Inches(12.25), Inches(7.05), Inches(.65), Inches(.25),
        [[(str(PAGE["n"]), 9, GREY, False)]], PP_ALIGN.RIGHT)


def head(s, title, kicker=None, color=BLUE):
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, Inches(.28), Inches(1.55), color)
    if kicker:
        txt(s, Inches(.85), Inches(.48), Inches(11.6), Inches(.35), [[(kicker, 13, color, True)]])
    # Keep the title clear of long kicker text.  The assessment kicker is wider
    # than the standard topic label, so its title starts on a separate baseline.
    title_y = Inches(1.05) if kicker and len(kicker) > 44 else Inches(.88)
    if len(title) > 54 and " — " in title:
        first, second = title.split(" — ", 1)
        txt(s, Inches(.85), title_y - Inches(.09), Inches(11.8), Inches(.86),
            [[(first + " —", 20, INK, True)], [(second, 20, INK, True)]], space=0)
    else:
        title_size = 23 if len(title) > 46 else 29
        txt(s, Inches(.85), title_y, Inches(11.8), Inches(.62), [[(title, title_size, INK, True)]])
    rect(s, Inches(.85), Inches(1.7), Inches(11.63), Inches(.02), LINE)
    return s


def cover():
    s = slide()
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, SW, Inches(.22), BLUE)
    rect(s, 0, Inches(7.28), SW, Inches(.22), TEAL)
    logo = os.path.join(ASSETS, "tertiary-infotech-logo.png")
    if os.path.exists(logo):
        s.shapes.add_picture(logo, Inches(.85), Inches(.7), height=Inches(1.02))
    rect(s, Inches(11.0), Inches(.72), Inches(1.55), Inches(1.0), DEEP)
    txt(s, Inches(11.0), Inches(.82), Inches(1.55), Inches(.45), [[("FMEA", 20, WHITE, True)]], PP_ALIGN.CENTER)
    txt(s, Inches(11.0), Inches(1.28), Inches(1.55), Inches(.25), [[("PREVENT RISK", 8, WHITE, True)]], PP_ALIGN.CENTER)
    txt(s, Inches(.9), Inches(2.25), Inches(11.6), Inches(.4), [[("COURSE SLIDES  ·  WSQ", 16, BLUE, True)]])
    txt(s, Inches(.9), Inches(2.82), Inches(11.6), Inches(1.7), [[(C.TITLE, 40, INK, True)]])
    rect(s, Inches(.92), Inches(4.72), Inches(2.4), Inches(.06), TEAL)
    txt(s, Inches(.9), Inches(5.02), Inches(11.8), Inches(1.1),
        [[(f"WSQ Course Code: {C.COURSE_CODE}", 16, GREY, False)],
         [("Conducted by Tertiary Infotech Academy Pte Ltd  ·  UEN 201200696W", 14, GREY, False)]])
    txt(s, Inches(.9), Inches(6.45), Inches(11.8), Inches(.3), [[(f"Version {C.VERSION}  ·  {C.VERSION_DATE}", 12, GREY, False)]])
    txt(s, Inches(.9), Inches(6.84), Inches(11.8), Inches(.25),
        [[("© 2026 Tertiary Infotech Academy Pte Ltd. All rights reserved.", 10, GREY, False)]])


def section(kicker, title, number, sub=""):
    s = slide()
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, Inches(.28), SH, BLUE)
    rect(s, Inches(.85), Inches(2.45), Inches(.14), Inches(2.1), TEAL)
    txt(s, Inches(1.25), Inches(2.5), Inches(10.8), Inches(.4), [[(kicker, 18, BLUE, True)]])
    txt(s, Inches(1.25), Inches(3.0), Inches(10.9), Inches(1.2), [[(title, 39, INK, True)]])
    if sub:
        txt(s, Inches(1.27), Inches(4.45), Inches(10.7), Inches(.6), [[(sub, 16, GREY, False)]])
    txt(s, Inches(10), Inches(.6), Inches(2.7), Inches(1.4), [[(number, 72, LINE, True)]], PP_ALIGN.RIGHT)
    footer(s)


def content(title, items, kicker=None, size=18):
    s = head(slide(), title, kicker)
    bullets(s, Inches(.95), Inches(2.05), Inches(11.3), Inches(4.55), items, size=size, gap=12)
    footer(s)
    return s


def two_col(title, left, right, kicker=None, lhead="Perspective A", rhead="Perspective B", lsize=16, rsize=16):
    s = head(slide(), title, kicker)
    for x, color, label, items, size in [
        (Inches(.85), BLUE, lhead, left, lsize), (Inches(6.95), TEAL, rhead, right, rsize)
    ]:
        rect(s, x, Inches(1.95), Inches(5.55), Inches(4.7), LIGHT)
        rect(s, x, Inches(1.95), Inches(5.55), Inches(.11), color)
        txt(s, x + Inches(.28), Inches(2.2), Inches(4.95), Inches(.4), [[(label, 16, color, True)]])
        bullets(s, x + Inches(.28), Inches(2.78), Inches(4.95), Inches(3.45), items, size=size)
    footer(s)
    return s


def cards3(title, cards, kicker):
    s = head(slide(), title, kicker)
    for i, (color, label, items) in enumerate(cards[:3]):
        x = Inches(.85 + i * 4.15)
        rect(s, x, Inches(1.95), Inches(3.65), Inches(4.7), LIGHT)
        rect(s, x, Inches(1.95), Inches(3.65), Inches(.12), color)
        txt(s, x + Inches(.25), Inches(2.23), Inches(3.15), Inches(.55), [[(label, 19, color, True)]])
        bullets(s, x + Inches(.25), Inches(3.0), Inches(3.12), Inches(3.2), items, size=14)
    footer(s)
    return s


def big_statement(line1, line2, kicker, color=BLUE):
    s = slide()
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, Inches(.28), SH, color)
    txt(s, Inches(1.1), Inches(2.12), Inches(11), Inches(.4), [[(kicker, 15, color, True)]])
    txt(s, Inches(1.1), Inches(2.72), Inches(11.2), Inches(1.65), [[(line1, 37, INK, True)]])
    txt(s, Inches(1.12), Inches(4.78), Inches(10.9), Inches(1.05), [[(line2, 20, GREY, False)]])
    footer(s)
    return s


def tile_grid(title, items, kicker=None, cols=2, size=15, accent=BLUE):
    s = head(slide(), title, kicker, accent)
    rows = math.ceil(len(items) / cols)
    x0, y0, total_w, total_h = Inches(.85), Inches(1.95), Inches(11.63), Inches(4.78)
    gx, gy = Inches(.3), Inches(.22)
    cw = int((total_w - gx * (cols - 1)) / cols)
    ch = int((total_h - gy * (rows - 1)) / rows)
    for i, item in enumerate(items):
        row, col = divmod(i, cols)
        x, y = int(x0 + col * (cw + gx)), int(y0 + row * (ch + gy))
        color = PALETTE[i % len(PALETTE)]
        rect(s, x, y, cw, ch, LIGHT)
        rect(s, x, y, Inches(.1), ch, color)
        oval(s, x + Inches(.27), int(y + ch / 2 - Inches(.3)), Inches(.6), Inches(.6), color)
        txt(s, x + Inches(.27), int(y + ch / 2 - Inches(.3)), Inches(.6), Inches(.6),
            [[(str(i + 1), 18, WHITE, True)]], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        label, body = item if isinstance(item, tuple) else (str(item), "")
        lines = [[(label, size + 1, INK, True)]]
        if body:
            lines.append([(body, size - 2, GREY, False)])
        txt(s, x + Inches(1.08), y + Inches(.1), cw - Inches(1.3), ch - Inches(.16), lines,
            anchor=MSO_ANCHOR.MIDDLE, space=3)
    footer(s)
    return s


def flow_h(title, steps, kicker=None, color=BLUE):
    s = head(slide(), title, kicker, color)
    n = len(steps)
    gap, total_w = Inches(.3), Inches(11.63)
    cw = int((total_w - gap * (n - 1)) / n)
    for i, step in enumerate(steps):
        x = int(Inches(.85) + i * (cw + gap))
        rect(s, x, Inches(2.45), cw, Inches(3.2), LIGHT)
        rect(s, x, Inches(2.45), cw, Inches(.1), color)
        oval(s, int(x + cw / 2 - Inches(.4)), Inches(2.85), Inches(.8), Inches(.8), color)
        txt(s, int(x + cw / 2 - Inches(.4)), Inches(2.85), Inches(.8), Inches(.8),
            [[(str(i + 1), 27, WHITE, True)]], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(s, x + Inches(.14), Inches(4.03), cw - Inches(.28), Inches(1.25),
            [[(str(step), 13, INK, False)]], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < n - 1:
            txt(s, x + cw - Inches(.03), Inches(3.68), gap + Inches(.08), Inches(.45),
                [[("▶", 14, color, True)]], PP_ALIGN.CENTER)
    footer(s)
    return s


def trainer_slide(kicker, name, role, rows, initials, accent=BLUE):
    s = head(slide(), "About the Trainer", kicker, accent)
    rect(s, Inches(.85), Inches(1.95), Inches(3.65), Inches(4.7), LIGHT)
    oval(s, Inches(1.82), Inches(2.45), Inches(1.7), Inches(1.7), accent)
    txt(s, Inches(1.82), Inches(2.45), Inches(1.7), Inches(1.7), [[(initials, 43, WHITE, True)]],
        PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    txt(s, Inches(1.05), Inches(4.55), Inches(3.25), Inches(.55), [[(name, 20, INK, True)]], PP_ALIGN.CENTER)
    txt(s, Inches(1.05), Inches(5.2), Inches(3.25), Inches(.9), [[(role, 13, GREY, False)]], PP_ALIGN.CENTER)
    tile_grid_rows = len(rows)
    h = int((Inches(4.7) - Inches(.18) * (tile_grid_rows - 1)) / tile_grid_rows)
    for i, (label, value) in enumerate(rows):
        y = int(Inches(1.95) + i * (h + Inches(.18)))
        rect(s, Inches(4.9), y, Inches(7.58), h, LIGHT)
        rect(s, Inches(4.9), y, Inches(.1), h, PALETTE[i % 4])
        shown = value or "____________________________________________"
        txt(s, Inches(5.22), y, Inches(6.95), h,
            [[(label.upper(), 10, PALETTE[i % 4], True)], [(shown, 13, INK if value else LINE, False)]],
            anchor=MSO_ANCHOR.MIDDLE, space=2)
    footer(s)
    return s


def activity_overview(activity):
    s = head(slide(), activity["title"], f"ACTIVITY {activity['num']:02d} · {activity['codes']}", TEAL)
    rect(s, Inches(.85), Inches(2.0), Inches(7.45), Inches(4.55), LIGHT)
    txt(s, Inches(1.15), Inches(2.32), Inches(6.85), Inches(.4), [[("ORION SCENARIO", 13, TEAL, True)]])
    txt(s, Inches(1.15), Inches(2.92), Inches(6.85), Inches(1.35), [[(activity["scenario"], 18, INK, False)]])
    txt(s, Inches(1.15), Inches(4.7), Inches(6.85), Inches(.35), [[("YOUR MISSION", 13, BLUE, True)]])
    txt(s, Inches(1.15), Inches(5.16), Inches(6.85), Inches(.95), [[(activity["desc"], 17, INK, True)]])
    rect(s, Inches(8.62), Inches(2.0), Inches(3.86), Inches(2.05), DEEP)
    txt(s, Inches(8.92), Inches(2.32), Inches(3.25), Inches(.35), [[("DELIVERABLE", 12, TEAL, True)]])
    txt(s, Inches(8.92), Inches(2.88), Inches(3.25), Inches(.92), [[(activity["build"], 15, WHITE, True)]])
    rect(s, Inches(8.62), Inches(4.3), Inches(3.86), Inches(2.25), LIGHT)
    txt(s, Inches(8.92), Inches(4.62), Inches(3.25), Inches(.35), [[("TOOLS AND FILES", 12, VIOLET, True)]])
    txt(s, Inches(8.92), Inches(5.15), Inches(3.25), Inches(1.0), [[(activity["services"], 14, INK, False)]])
    footer(s)


def activity_target(activity):
    tile_grid(f"Activity {activity['num']:02d} — Decision Target", [
        ("Objective", activity["objective"]),
        ("Question 1", activity["questions"][0]),
        ("Question 2", activity["questions"][1]),
        ("Evidence standard", activity["test"]),
    ], kicker="WHAT GOOD WORK LOOKS LIKE · PROCEDURE IN LEARNER GUIDE", cols=2, size=14, accent=TEAL)


def test_slide(activity):
    s = head(slide(), f"Activity {activity['num']:02d} — Learning Check", "REVIEW BEFORE SUBMISSION", TEAL)
    rect(s, Inches(.85), Inches(2.02), Inches(7.35), Inches(4.5), RGBColor(0xE8, 0xF7, 0xEE))
    txt(s, Inches(1.18), Inches(2.35), Inches(6.7), Inches(.42), [[("✓  Acceptance check", 20, DEEP, True)]])
    txt(s, Inches(1.18), Inches(3.08), Inches(6.55), Inches(1.7), [[(activity["test"], 18, INK, False)]])
    txt(s, Inches(1.18), Inches(5.38), Inches(6.5), Inches(.55),
        [[("PASS = complete · traceable · evidence-based · saved", 14, DEEP, True)]])
    rect(s, Inches(8.5), Inches(2.02), Inches(3.98), Inches(2.05), LIGHT)
    txt(s, Inches(8.82), Inches(2.34), Inches(3.3), Inches(.4), [[("EVIDENCE", 13, BLUE, True)]])
    bullets(s, Inches(8.82), Inches(2.9), Inches(3.15), Inches(.95),
            ["Completed workbook", "Recorded decision and rationale"], size=13, gap=7)
    rect(s, Inches(8.5), Inches(4.34), Inches(3.98), Inches(2.18), LIGHT)
    txt(s, Inches(8.82), Inches(4.65), Inches(3.3), Inches(.4), [[("SELF-CHECK", 13, AMBER, True)]])
    bullets(s, Inches(8.82), Inches(5.2), Inches(3.15), Inches(1.05),
            ["Formula outputs reconcile", "Assumptions are visible", "Next action has an owner"], size=12, gap=6)
    footer(s)


def step_slide(kicker, title, n, total, text, note=""):
    """Used as a method spotlight—not as application click-by-click guidance."""
    s = head(slide(), title, kicker, VIOLET)
    oval(s, Inches(.9), Inches(2.45), Inches(1.45), Inches(1.45), VIOLET)
    txt(s, Inches(.9), Inches(2.45), Inches(1.45), Inches(1.45), [[(str(n), 40, WHITE, True)]],
        PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    txt(s, Inches(.98), Inches(1.98), Inches(4), Inches(.3), [[(f"METHOD STAGE {n} OF {total}", 12, GREY, True)]])
    txt(s, Inches(2.7), Inches(2.35), Inches(9.55), Inches(1.45), [[(text, 24, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    if note:
        rect(s, Inches(2.7), Inches(4.25), Inches(9.55), Inches(1.25), LIGHT)
        txt(s, Inches(3.02), Inches(4.52), Inches(8.95), Inches(.72), [[(note, 16, GREY, False)]])
    footer(s)
    return s


def brk(kind, subtitle, color=AMBER):
    s = slide()
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, SW, Inches(.22), color)
    rect(s, 0, Inches(7.28), SW, Inches(.22), color)
    rect(s, Inches(5.4), Inches(2.35), Inches(2.53), Inches(.1), color)
    txt(s, 0, Inches(2.75), SW, Inches(1.0), [[(kind, 46, INK, True)]], PP_ALIGN.CENTER)
    txt(s, 0, Inches(4.0), SW, Inches(.75), [[(subtitle, 21, color, True)]], PP_ALIGN.CENTER)
    PAGE["n"] += 1


def special_visual(kind, title, payload, kicker):
    if kind == "pdca":
        colors = [BLUE, TEAL, VIOLET, AMBER]
        cards = [(colors[i], item.split(":", 1)[0], [item.split(":", 1)[1].strip()]) for i, item in enumerate(payload)]
        s = head(slide(), title, kicker, TEAL)
        positions = [(1.2, 2.0), (7.0, 2.0), (7.0, 4.45), (1.2, 4.45)]
        for i, (color, label, body) in enumerate(cards):
            x, y = positions[i]
            rect(s, Inches(x), Inches(y), Inches(5.0), Inches(1.55), LIGHT)
            rect(s, Inches(x), Inches(y), Inches(.12), Inches(1.55), color)
            oval(s, Inches(x + .3), Inches(y + .38), Inches(.72), Inches(.72), color)
            txt(s, Inches(x + .3), Inches(y + .38), Inches(.72), Inches(.72), [[(str(i + 1), 22, WHITE, True)]], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            txt(s, Inches(x + 1.25), Inches(y + .22), Inches(3.45), Inches(1.1), [[(label, 16, color, True)], [(body[0], 13, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
        txt(s, Inches(5.64), Inches(3.7), Inches(2.05), Inches(.45), [[("LEARN ↻", 17, TEAL, True)]], PP_ALIGN.CENTER)
        footer(s)
        return
    if kind == "pareto":
        import matplotlib.pyplot as plt
        values, labels = payload, C.CONTENT_SLIDES[2][3][3]
        cum, running, total = [], 0, sum(values)
        for value in values:
            running += value
            cum.append(running / total * 100)
        fig, ax = plt.subplots(figsize=(9, 4.5), dpi=160)
        ax.bar(labels, values, color="#1F6FEB")
        ax.set_ylabel("Defect count")
        ax.tick_params(axis="x", rotation=22, labelsize=9)
        ax.grid(axis="y", alpha=.18)
        ax2 = ax.twinx()
        ax2.plot(labels, cum, color="#F59E0B", marker="o", linewidth=2.5)
        ax2.set_ylim(0, 110)
        ax2.set_ylabel("Cumulative %")
        for spine in list(ax.spines.values()) + list(ax2.spines.values()):
            spine.set_visible(False)
        fig.tight_layout()
        image = os.path.join(tempfile.gettempdir(), "fmea-pareto.png")
        fig.savefig(image, facecolor="white", bbox_inches="tight")
        plt.close(fig)
        s = head(slide(), title, kicker)
        s.shapes.add_picture(image, Inches(.82), Inches(1.98), width=Inches(8.0), height=Inches(4.65))
        rect(s, Inches(9.08), Inches(2.0), Inches(3.38), Inches(4.6), LIGHT)
        txt(s, Inches(9.42), Inches(2.36), Inches(2.7), Inches(.35), [[("READ THE SIGNAL", 13, AMBER, True)]])
        bullets(s, Inches(9.42), Inches(2.95), Inches(2.65), Inches(2.7),
                ["Solder bridge is the largest category.", "The cumulative line shows concentration.", "Priority is not proof of cause."], size=14)
        footer(s)
        return
    if kind == "fishbone":
        s = head(slide(), title, kicker, TEAL)
        rect(s, Inches(2.0), Inches(4.0), Inches(8.7), Inches(.06), TEAL)
        txt(s, Inches(10.75), Inches(3.58), Inches(1.6), Inches(.85), [[("BRIDGE\nDEFECT", 16, RED, True)]], PP_ALIGN.CENTER)
        for i, label in enumerate(payload):
            upper = i < 3
            x = 2.15 + (i % 3) * 2.65
            y = 2.2 if upper else 4.65
            rect(s, Inches(x), Inches(y), Inches(2.15), Inches(.72), LIGHT)
            txt(s, Inches(x), Inches(y + .15), Inches(2.15), Inches(.35), [[(label, 14, PALETTE[i % 4], True)]], PP_ALIGN.CENTER)
            line = s.shapes.add_shape(1, Inches(x + .95), Inches(2.9 if upper else 4.03), Inches(.05), Inches(1.12 if upper else .62))
            line.fill.solid(); line.fill.fore_color.rgb = PALETTE[i % 4]; line.line.fill.background()
        footer(s)
        return
    if kind == "faulttree":
        s = head(slide(), title, kicker, VIOLET)
        rect(s, Inches(4.75), Inches(2.0), Inches(3.75), Inches(.85), RED)
        txt(s, Inches(4.75), Inches(2.18), Inches(3.75), Inches(.4), [[(payload[0], 17, WHITE, True)]], PP_ALIGN.CENTER)
        rect(s, Inches(6.60), Inches(2.84), Inches(.05), Inches(.33), VIOLET)
        oval(s, Inches(6.05), Inches(3.15), Inches(1.15), Inches(1.15), VIOLET)
        txt(s, Inches(6.05), Inches(3.15), Inches(1.15), Inches(1.15), [[("OR", 18, WHITE, True)]], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        rect(s, Inches(6.60), Inches(4.28), Inches(.05), Inches(.31), VIOLET)
        rect(s, Inches(2.76), Inches(4.56), Inches(8.32), Inches(.05), VIOLET)
        for i, label in enumerate(["Excess paste deposited", "Components bridge after placement", "Paste slumps during reflow"]):
            x = 1.0 + i * 4.15
            rect(s, Inches(x + 1.75), Inches(4.56), Inches(.05), Inches(.22), VIOLET)
            rect(s, Inches(x), Inches(4.78), Inches(3.55), Inches(1.15), LIGHT)
            txt(s, Inches(x + .2), Inches(4.98), Inches(3.15), Inches(.7), [[(label, 14, INK, True)]], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(s, Inches(1.15), Inches(6.25), Inches(11), Inches(.35), [[("Gates make causal combinations explicit; evidence tests the basic events.", 14, VIOLET, True)]], PP_ALIGN.CENTER)
        footer(s)
        return
    if kind in {"failurechain", "controlplan"}:
        flow_h(title, payload, kicker=kicker, color=TEAL if kind == "failurechain" else VIOLET)
        return
    if kind == "riskmatrix":
        cards3(title, [
            (RED, "SEVERITY", ["Consequence of the effect", "Score the most serious credible effect"]),
            (AMBER, "OCCURRENCE", ["Likelihood of the stated cause", "Use prevention and evidence"]),
            (VIOLET, "DETECTION", ["Ability to detect before impact", "Higher score means weaker detection"]),
        ], kicker)
        return
    if kind == "controlchart":
        import matplotlib.pyplot as plt
        values = [99.7, 100.5, 100.1, 99.8, 100.4, 100.2, 99.9, 100.3, 100.1, 99.6,
                  100.4, 100.6, 100.2, 100.0, 99.7, 100.3, 100.5, 100.1, 99.9, 100.2,
                  100.4, 100.7, 101.5, 100.6, 100.3, 100.2, 100.1, 99.8, 100.0, 100.3]
        cl, ucl, lcl = 100.15, 101.05, 99.25
        fig, ax = plt.subplots(figsize=(9, 4.4), dpi=160)
        ax.plot(range(1, 31), values, color="#1F6FEB", marker="o", linewidth=1.8)
        ax.axhline(cl, color="#10B981", linewidth=2, label="CL")
        ax.axhline(ucl, color="#DC2626", linestyle="--", label="UCL")
        ax.axhline(lcl, color="#DC2626", linestyle="--", label="LCL")
        ax.scatter([23], [101.5], s=95, color="#DC2626", zorder=5)
        ax.set_xlabel("Time order")
        ax.set_ylabel("Solder-paste volume index")
        ax.grid(alpha=.16)
        ax.legend(frameon=False, ncol=3)
        for spine in ax.spines.values(): spine.set_visible(False)
        fig.tight_layout()
        image = os.path.join(tempfile.gettempdir(), "fmea-control-chart.png")
        fig.savefig(image, facecolor="white", bbox_inches="tight")
        plt.close(fig)
        s = head(slide(), title, kicker, VIOLET)
        s.shapes.add_picture(image, Inches(.75), Inches(2.0), width=Inches(8.2), height=Inches(4.55))
        rect(s, Inches(9.2), Inches(2.0), Inches(3.25), Inches(4.55), LIGHT)
        txt(s, Inches(9.53), Inches(2.35), Inches(2.6), Inches(.38), [[("CONTROL ≠ SPECIFICATION", 12, VIOLET, True)]])
        bullets(s, Inches(9.53), Inches(2.95), Inches(2.55), Inches(2.9),
                ["Limits describe process behaviour.", "A signal prompts investigation.", "Capability compares stable variation with specifications."], size=13)
        footer(s)
        return


def activities_access():
    s = head(slide(), "Access the Hands-On Activities", "COURSE REPOSITORY", TEAL)
    cards3_data = [
        (BLUE, "Download", ["Open the GitHub repository", "Choose Code → Download ZIP", "Extract before editing files"]),
        (TEAL, "Work", ["One folder per activity", "Use the starter workbook", "Follow detailed steps in the Learner Guide"]),
        (VIOLET, "Verify", ["Run the folder checklist", "Compare with the completed sample", "Save your evidence clearly"]),
    ]
    for i, card in enumerate(cards3_data):
        x = Inches(.85 + i * 4.15)
        rect(s, x, Inches(1.95), Inches(3.65), Inches(3.9), LIGHT)
        rect(s, x, Inches(1.95), Inches(3.65), Inches(.12), card[0])
        txt(s, x + Inches(.25), Inches(2.22), Inches(3.15), Inches(.45), [[(card[1], 19, card[0], True)]])
        bullets(s, x + Inches(.25), Inches(2.95), Inches(3.1), Inches(2.5), card[2], size=13)
    box = s.shapes.add_textbox(Inches(.95), Inches(6.12), Inches(11.4), Inches(.42))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = C.GITHUB_URL
    run.font.name = "Arial"; run.font.size = Pt(14); run.font.bold = True; run.font.color.rgb = BLUE
    run.hyperlink.address = C.GITHUB_URL
    footer(s)


def lesson_plan_slide():
    tile_grid("Two-Day Learning Journey", [
        ("Day 1 · AM", "Administration, Topic 1 and Activities 1–2: CI system, SIPOC and PDCA."),
        ("Day 1 · PM", "Topic 2 and Activities 3–4: Pareto, Fishbone, 5 Whys and Fault Tree."),
        ("Day 2 · AM", "Topic 3 and Activities 5–6: build, rate and optimise the Orion PFMEA."),
        ("Day 2 · PM", "Topics 4–5, Activities 7–10, then 60-minute WA and 90-minute PP."),
    ], kicker="16 HOURS · 13.5 HOURS TRAINING + 2.5 HOURS ASSESSMENT", cols=2, size=14, accent=TEAL)


def course_material_slide():
    flow_h("Download Course Materials", ["Open LMS/TMS", "Select this course", "Open courseware links", "Download and extract files", "Open the matching activity folder"],
           kicker="LMS-TMS.TERTIARYINFOTECH.COM · VISUAL DOWNLOAD PATH", color=BLUE)


def practice_exam_slide():
    s = head(slide(), "Practice and Readiness Check", "PRACTICE EXAM PORTAL", VIOLET)
    cards = [
        (BLUE, "Review", ["Revisit K1–K4 and A1–A10", "Open completed activity evidence", "Explain decisions, not only scores"]),
        (TEAL, "Practise", ["Use the assigned FMEA practice set", "Work individually first", "Retain formulas and rationale"]),
        (VIOLET, "Reflect", ["Trace action to failure mechanism", "Check residual risk and controls", "State one limitation and next step"]),
    ]
    for i, card in enumerate(cards):
        x = Inches(.85 + i * 4.15)
        rect(s, x, Inches(1.95), Inches(3.65), Inches(3.9), LIGHT)
        rect(s, x, Inches(1.95), Inches(3.65), Inches(.12), card[0])
        txt(s, x + Inches(.25), Inches(2.22), Inches(3.15), Inches(.45), [[(card[1], 19, card[0], True)]])
        bullets(s, x + Inches(.25), Inches(2.95), Inches(3.1), Inches(2.5), card[2], size=13)
    url = "https://exams.tertiaryinfotech.com/"
    box = s.shapes.add_textbox(Inches(.95), Inches(6.12), Inches(11.4), Inches(.42))
    p = box.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = url; run.font.name = "Arial"; run.font.size = Pt(14); run.font.bold = True; run.font.color.rgb = VIOLET
    run.hyperlink.address = url
    footer(s)


def sources_slide():
    s = head(slide(), "Key FMEA References", "SOURCE-ALIGNED COURSEWARE", VIOLET)
    sources = [
        ("ASQ", "FMEA purpose, timing and team method", "https://asq.org/quality-resources/fmea"),
        ("Jama Software", "Seven-step method, Action Priority and RPN cautions", "https://www.jamasoftware.com/requirements-management-guide/meeting-regulatory-compliance-and-industry-standards/fmea/"),
        ("Quality-One", "Functions, controls, action ownership and re-rating", "https://quality-one.com/fmea/"),
        ("PMC", "FMEA as a prospective safety and risk-control tool", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10229026/"),
        ("IPH Hannover", "German industrial FMEA practice and prevention context", "https://www.iph-hannover.de/de/dienstleistungen/fertigungsverfahren/fmea/"),
        ("ScienceDirect", "Engineering overview and related reliability concepts", "https://www.sciencedirect.com/topics/engineering/failure-modes-and-effects-analysis"),
    ]
    for i, (label, body, url) in enumerate(sources):
        row, col = divmod(i, 2)
        x, y = Inches(.85 + col * 6.1), Inches(1.95 + row * 1.48)
        rect(s, x, y, Inches(5.55), Inches(1.18), LIGHT)
        rect(s, x, y, Inches(.1), Inches(1.18), PALETTE[i % 4])
        box = s.shapes.add_textbox(x + Inches(.28), y + Inches(.16), Inches(1.8), Inches(.35))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = label; run.font.name = "Arial"; run.font.size = Pt(12); run.font.bold = True
        run.font.color.rgb = PALETTE[i % 4]; run.hyperlink.address = url
        txt(s, x + Inches(2.08), y + Inches(.12), Inches(3.18), Inches(.82), [[(body, 11, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s)


def assessment_flow():
    flow_h("Assessment Flow", ["Assessment attendance", "Written Assessment · K1–K4 · 60 min", "Practical Performance · A1–A10 · 90 min", "Submit candidate evidence", "Sign Assessment Summary Record"],
           kicker="INDIVIDUAL · OPEN BOOK · COMPETENT / NOT YET COMPETENT", color=VIOLET)


SLIDE_MAP = {"topics": {}, "activities": {}}
cover()
section("COURSE ADMINISTRATION", "Welcome & Housekeeping", "")
tile_grid("Digital Attendance (Mandatory)", [
    ("AM attendance", "Scan the SSG digital-attendance QR code when class starts."),
    ("PM attendance", "Scan again after lunch when instructed."),
    ("Assessment attendance", "Complete the separate attendance before assessment."),
    ("Funding requirement", "At least 75% attendance and a Competent result are required."),
], kicker="TRAQOM · SSG", cols=2, accent=TEAL)
trainer_slide("YOUR TRAINER · GENERAL TEMPLATE", "Trainer Profile Template", "Trainer completes this profile before delivery",
              [("Name", "[Trainer name]"), ("Designation", "[Title / role]"), ("Qualifications", "[Relevant qualifications]"),
               ("Expertise", "[Subject-matter strengths]"), ("Experience", "[Industry and training experience]"), ("Contact", "[Approved contact channel]")], "T", GREY)
trainer_slide("YOUR TRAINER", C.TRAINER, "Principal Trainer · Tertiary Infotech Academy Pte Ltd",
              [("Role", "Principal Trainer and courseware developer"), ("Expertise", "Process improvement, analytics, automation and AI"),
               ("Delivers", "WSQ data, software, cloud and AI programmes"), ("Founder", "Tertiary Infotech Academy / Tertiary Courses")], "AA", TEAL)
cards3("Let's Know Each Other", [
    (BLUE, "You", ["Name and organisation", "Role and process context", "One improvement challenge"]),
    (TEAL, "Your process", ["Customer and output", "One recurring failure", "Current measure"]),
    (VIOLET, "Your goal", ["Risk to prevent", "Capability to improve", "Learning outcome for these two days"]),
], "INTRODUCTION")
tile_grid("Ground Rules", ["Set mobile phones to silent mode.", "Participate actively—no question is too small.",
          "Respect different views; agree to disagree.", "One conversation at a time.",
          "Return from breaks on time.", "Maintain at least 75% attendance."],
          kicker="HOUSEKEEPING", cols=2, size=15)
tile_grid("Course Outline", [(f"Topic {t['num']}", t["title"]) for t in C.TOPICS],
          kicker="FIVE CONNECTED TOPICS · ONE ORION CASE", cols=2, size=14, accent=BLUE)
lesson_plan_slide()
tile_grid("Learning Outcomes", [(f"ELO{i}", value.split(":", 1)[1].strip()) for i, value in enumerate(C.LEARNING_OUTCOMES, 1)],
          kicker="BY THE END OF THE COURSE", cols=2, size=14, accent=TEAL)
two_col("Skills Framework Alignment", C.KNOWLEDGE, C.ABILITIES,
        kicker=f"{C.TSC_TITLE} · {C.TSC_CODE}", lhead="Knowledge K1–K4", rhead="Abilities A1–A10", lsize=14, rsize=11)
tile_grid("Briefing for Assessment", [
    ("Individual work", "No discussion or copying during assessment."),
    ("Open book", "Use course slides, Learner Guide and approved materials."),
    ("Evidence", "Complete every response and practical deliverable."),
    ("Conduct", "No photos; phones and unrelated materials remain away."),
    ("Writing", "Use black/blue pen where hard-copy responses are required."),
    ("Rights", "Re-assessment and appeal processes are available."),
], kicker="BEFORE ASSESSMENT", cols=2, size=13, accent=AMBER)
cards3("Final Assessment", [
    (BLUE, "Written Assessment", ["4 open-ended questions", "K1–K4", "60 minutes"]),
    (TEAL, "Practical Performance", ["5 connected tasks", "A1–A10", "90 minutes"]),
    (VIOLET, "Competency", ["Open book", "C / NYC outcome", "Submit evidence and sign record"]),
], "WA + PP · 150 MINUTES")
assessment_flow()
activities_access()
course_material_slide()
practice_exam_slide()
tile_grid("Connected Case — Orion Precision Electronics", [
    ("Process", C.CASE_STUDY["process"]), ("Signal", C.CASE_STUDY["problem"]),
    ("Goal", C.CASE_STUDY["goal"]), ("Learning path", "Define → analyse → anticipate → improve → control"),
], kicker="ONE CASE · TEN ACTIVITIES", cols=2, size=13, accent=DEEP)

topic_activities = {t["num"]: [a for a in ACTIVITIES if a["topic"] == t["num"]] for t in C.TOPICS}
for topic in C.TOPICS:
    if topic["num"] == 3:
        brk("Day 2", "FMEA design · improvement action · process control")
    section(f"TOPIC {topic['code']}", topic["title"], topic["code"], topic["subtitle"])
    SLIDE_MAP["topics"][str(topic["num"])] = {"divider": len(prs.slides), "content": []}
    accent = PALETTE[(topic["num"] - 1) % 4]
    key_titles = {
        1: "Key Concepts — Continuous Improvement",
        2: "Key Concepts — Root Cause Analysis",
        3: "Key Concepts — Design and Apply FMEA",
        4: "Key Concepts — Improvement Opportunities",
        5: "Key Concepts — Process Control Plan",
    }
    tile_grid(key_titles[topic["num"]], topic["concepts"], kicker=topic["weighting"], cols=2, size=13, accent=accent)
    for spec in C.CONTENT_SLIDES[topic["num"]]:
        kind, title, payload = spec[0], spec[1], spec[2]
        kicker = f"TOPIC {topic['code']}"
        if kind == "big":
            big_statement(title, payload, f"{kicker} · WHY IT MATTERS", accent)
        elif kind == "tiles":
            tile_grid(title, payload, kicker=kicker, cols=2, size=13, accent=accent)
        elif kind == "flow":
            flow_h(title, payload, kicker=kicker, color=accent)
        elif kind == "two":
            two_col(title, payload, spec[3], kicker=kicker, lhead="Use", rhead="Contrast / check", lsize=15, rsize=15)
        elif kind == "cards":
            cards3(title, [(PALETTE[i % 4], item[0], item[1]) for i, item in enumerate(payload)], kicker)
        else:
            special_visual(kind, title, payload, kicker)
        SLIDE_MAP["topics"][str(topic["num"])]["content"].append({"title": title, "slide": len(prs.slides), "kind": kind})
        if topic["num"] == 3 and title == "Seven-step FMEA method":
            step_slide("TOPIC 03 · METHOD SPOTLIGHT", "Failure Analysis", 4, 7,
                       "Connect function, failure mode, effect and cause before assigning scores.",
                       "The exact workbook procedure is in Activity 05 and the Learner Guide.")
    for activity in topic_activities[topic["num"]]:
        start = len(prs.slides) + 1
        activity_overview(activity)
        activity_target(activity)
        test_slide(activity)
        SLIDE_MAP["activities"][str(activity["num"])] = {"title": activity["title"], "start": start, "end": len(prs.slides)}
    tile_grid(f"Recap — {topic['title']}", [
        ("Concept", topic["concepts"][0][1]),
        ("Decision", topic["concepts"][1][1]),
        ("Application", topic_activities[topic["num"]][0]["build"]),
        ("Transfer", "Name one workplace risk, evidence source and next action."),
    ], kicker="TOPIC RECAP", cols=2, size=13, accent=TEAL)
    SLIDE_MAP["topics"][str(topic["num"])]["end"] = len(prs.slides)

section("WRAP-UP", "Course Summary & Assessment", "")
tile_grid("What You Achieved", [(f"ELO{i}", value.split(":", 1)[1].strip()) for i, value in enumerate(C.LEARNING_OUTCOMES, 1)],
          kicker="LEARNING OUTCOMES", cols=2, size=14, accent=TEAL)
sources_slide()
tile_grid("Assessment Conditions", [
    ("Individual", "Complete both assessment instruments independently."),
    ("Approved resources", "Use only the course materials allowed by the assessor."),
    ("Evidence", "Submit the exact evidence named in each task."),
    ("Clarification", "Ask about instructions—not answers."),
], kicker="FINAL REMINDER", cols=2, size=14, accent=BLUE)
assessment_flow()
tile_grid("Digital Attendance (Mandatory)", [
    ("Assessment attendance", "Scan before the assessment starts."),
    ("Competency record", "Complete WA and PP and sign the summary record."),
    ("Submission", "Upload only the required learner evidence."),
    ("Support", "Ask the assessor about re-assessment or appeal if required."),
], kicker="TRAQOM · SSG", cols=2, size=14, accent=TEAL)
big_statement("Thank You!", "Prevent failure, improve the system and make learning visible.", "FROM RISK TO CONTROL", TEAL)

OUT = os.path.join(REPO, "courseware", f"{C.SHORT_TITLE}-{C.VERSION}.pptx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)


def add_fade_transitions(path):
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    prefix = f"{{{ns}}}"
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as temp:
        temp_path = temp.name
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", item.filename):
                from lxml import etree
                root = etree.fromstring(data)
                old = root.find(f"{prefix}transition")
                if old is not None:
                    root.remove(old)
                transition = etree.Element(f"{prefix}transition", spd="med", advClick="1")
                etree.SubElement(transition, f"{prefix}fade")
                c_sld = root.find(f"{prefix}cSld")
                root.insert(root.index(c_sld) + 1, transition)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(item, data)
    os.replace(temp_path, path)


add_fade_transitions(OUT)
with open(os.path.join(HERE, "slide_map.json"), "w", encoding="utf-8") as handle:
    json.dump(SLIDE_MAP, handle, indent=2)
print(f"Saved {OUT} ({len(prs.slides)} slides)")
