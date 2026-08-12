#!/usr/bin/env python3
"""Build the two-day Lesson Plan with exact references to the current deck."""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import course_data as C
import prodoc

ROOT = HERE.parents[1]
OUT = ROOT / "courseware"
LOGO = OUT / "assets" / "tertiary-infotech-logo.png"
BLUE = RGBColor(0x1F, 0x6F, 0xEB)


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), color); tcpr.append(shd)


def activity_ref(mapping, number):
    item = mapping["activities"][str(number)]
    return f"Slides {item['start']}–{item['end']}"


def topic_ref(mapping, number):
    item = mapping["topics"][str(number)]
    return f"Slides {item['divider']}–{item['end']}"


def add_schedule(doc, day, rows, total):
    doc.add_heading(f"Day {day} Schedule", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    labels = ["Time", "Duration", "Topic / activity", "Learning outcome and content", "Facilitation / learner activity", "Resources / slides", "Evidence / check"]
    for i, label in enumerate(labels):
        cell = table.rows[0].cells[i]; cell.text = label; shade(cell, "1F6FEB")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); run.bold = True; run.font.size = Pt(8)
    widths = [.85, .7, 1.2, 2.0, 2.1, 1.65, 1.8]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
        table.rows[0].cells[idx].width = Inches(width)
    for data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(data):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cells[idx].paragraphs:
                for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(8)
            cells[idx].width = Inches(widths[idx])
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    p = doc.add_paragraph(); r = p.add_run(f"Day {day} total: {total} hours"); r.bold = True; r.font.color.rgb = BLUE


def main():
    mapping = json.loads((HERE / "slide_map.json").read_text(encoding="utf-8"))
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Inches(.65); section.bottom_margin = Inches(.72); section.left_margin = Inches(.55); section.right_margin = Inches(.55)
        section.orientation = 1
        section.page_width, section.page_height = section.page_height, section.page_width
    prodoc.style_headings(doc)
    prodoc.add_cover_page(doc, "Lesson Plan", C.TITLE, C.VERSION, org_logo=str(LOGO), course_code=C.COURSE_CODE)
    prodoc.add_version_control(doc, [("v10.0", "Historical reference", "Legacy lesson structure", C.ORG),
                                     (C.VERSION, C.VERSION_DATE, "Aligned 2-day plan, ten activities and current 146-slide deck", C.ORG)])
    prodoc.add_toc(doc, "1-2")
    doc.add_heading("Course Information", level=1)
    table = doc.add_table(rows=0, cols=2); table.style = "Table Grid"
    for label, value in [("Course", C.TITLE), ("Course code", C.COURSE_CODE), ("TSC", f"{C.TSC_TITLE} ({C.TSC_CODE})"),
                         ("Duration", "2 days / 16 hours"), ("Training", "13.5 hours"),
                         ("Assessment", "2.5 hours: WA 60 minutes + PP 90 minutes"),
                         ("Method", "Facilitated concepts, case discussion, demonstration, ten connected hands-on activities and reflective review")]:
        cells = table.add_row().cells; cells[0].text = label; cells[1].text = value; cells[0].paragraphs[0].runs[0].bold = True
    doc.add_heading("Learning Outcomes", level=1)
    for item in C.LEARNING_OUTCOMES: doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()
    day1 = [
        ("09:00–10:30", "1.5 h", "Administration and course orientation", "Attendance, introductions, outcomes, competency mapping, assessment briefing and Orion case.", "Facilitator briefing; learner introduces a workplace process and improvement need.", "Slides 1–17", "Participation; assessment conditions understood"),
        ("10:30–10:45", "Included", "Short break", "Refreshment break within the scheduled session.", "Return on time.", "—", "Attendance maintained"),
        ("10:45–13:00", "2.25 h", "Topic 1 + Activities 01–02", "Create a CI system; culture, governance, PDCA, SIPOC, customer/process voice and coaching.", "Case analysis; complete CI charter and SIPOC/PDCA workbooks; peer scope review.", f"{topic_ref(mapping,1)}; {activity_ref(mapping,1)}; {activity_ref(mapping,2)}", "Charter, roles, governance, SIPOC and PDCA hypothesis"),
        ("13:00–14:00", "Not counted", "Lunch", "Lunch break.", "—", "—", "—"),
        ("14:00–15:30", "1.5 h", "Topic 2 concepts", "Problem statements, cause/effect levels, Pareto, Fishbone, evidence-backed 5 Whys, Fault Tree and data quality.", "Interpret a defect Pareto; distinguish symptom, mode, effect and cause; test causal statements.", topic_ref(mapping,2), "Bounded investigation statement and candidate-cause tests"),
        ("15:30–15:45", "Included", "Short break", "Refreshment break within the scheduled session.", "Return on time.", "—", "Attendance maintained"),
        ("15:45–18:00", "2.25 h", "Activities 03–04 + Day 1 review", "Prioritise the defect pattern and verify root causes using complementary tools.", "Complete Pareto, stratification, Fishbone, 5 Whys, Fault Tree and verification matrix; explain one confirmed mechanism.", f"{activity_ref(mapping,3)}; {activity_ref(mapping,4)}", "Reconciled analysis, verification evidence and Day 1 reflection"),
    ]
    add_schedule(doc, 1, day1, "8.0 scheduled training")
    doc.add_page_break()
    day2 = [
        ("09:00–10:30", "1.5 h", "Topic 3 concepts", "FMEA purpose, types, seven-step method, failure chain, controls, S/O/D, RPN cautions, action priority and review triggers.", "Construct and critique failure chains; compare equal-RPN risk profiles.", topic_ref(mapping,3), "Coherent function–mode–effect–cause logic"),
        ("10:30–10:45", "Included", "Short break", "Refreshment break within the scheduled session.", "Return on time.", "—", "Attendance maintained"),
        ("10:45–12:30", "1.75 h", "Activities 05–06", "Build, rate and optimise the Orion PFMEA.", "Complete PFMEA rows, S/O/D rationale, RPN formulas, actions, ownership, verification and re-score.", f"{activity_ref(mapping,5)}; {activity_ref(mapping,6)}", "Rated PFMEA and evidence-based residual-risk decision"),
        ("12:30–13:30", "Not counted", "Lunch", "Lunch break.", "—", "—", "—"),
        ("13:30–14:30", "1.0 h", "Topic 4 + Activities 07–08", "Select, pilot and direct risk-reduction actions with transparent criteria and coaching safeguards.", "Use weighted/Pugh comparison; draft a bounded pilot and coaching plan.", f"{topic_ref(mapping,4)}; {activity_ref(mapping,7)}; {activity_ref(mapping,8)}", "Recommendation and scale/revise/stop rule"),
        ("14:30–15:30", "1.0 h", "Topic 5 + Activities 09–10", "Variation, stability, capability, SPC signals, Control Plan, reaction plan and closed-loop review.", "Interpret a control chart; calculate capability; complete traceable operational controls.", f"{topic_ref(mapping,5)}; {activity_ref(mapping,9)}; {activity_ref(mapping,10)}", "Evaluation statement and deployable Control Plan"),
        ("15:30–16:30", "1.0 h", "Written Assessment (SAQ)", "Four open-ended questions covering K1–K4.", "Individual, open book under assessment conditions.", "WA paper; LMS link", "Completed WA responses"),
        ("16:30–18:00", "1.5 h", "Practical Performance", "Five connected tasks covering A1–A10.", "Individual completion of assigned practical evidence; submit and sign the Assessment Summary Record.", "PP paper; activity workbooks; LMS link", "Completed practical evidence and signed record"),
    ]
    add_schedule(doc, 2, day2, "5.5 training + 2.5 assessment = 8.0")
    doc.add_page_break(); doc.add_heading("Assessment Plan", level=1)
    table = doc.add_table(rows=1, cols=5); table.style = "Table Grid"
    for i, label in enumerate(["Instrument", "Duration", "Coverage", "Conditions", "Evidence"]):
        table.rows[0].cells[i].text = label; shade(table.rows[0].cells[i], "1F6FEB")
        for run in table.rows[0].cells[i].paragraphs[0].runs: run.font.color.rgb = RGBColor(255,255,255); run.bold = True
    for row in [("Written Assessment (SAQ)", "60 min", "K1–K4", "Individual; open book", "Four written responses"),
                ("Practical Performance", "90 min", "A1–A10", "Individual; use supplied files", "Five task outputs and assessor record")]:
        cells = table.add_row().cells
        for i, value in enumerate(row): cells[i].text = value
    doc.add_heading("Facilitator Preparation", level=1)
    for item in ["Verify digital-attendance access and learner materials before class.", "Download the current deck, Learner Guide and all ten activity folders.",
                 "Check Microsoft Excel can open, calculate and save the supplied files.", "Keep answer keys trainer-only and distribute only candidate papers at assessment.",
                 "After every activity, review the stated acceptance criterion and evidence—not only template completion."]:
        doc.add_paragraph(item, style="List Bullet")
    prodoc.add_page_numbers(doc, f"{C.TITLE} · {C.COURSE_CODE}")
    prodoc.enable_update_fields(doc)
    path = OUT / f"LP-{C.SHORT_TITLE}-{C.VERSION}.docx"
    doc.save(path)
    print(f"Saved {path}")


if __name__ == "__main__":
    main()
